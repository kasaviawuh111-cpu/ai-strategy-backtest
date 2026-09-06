from pathlib import Path
import json, html, shutil, sys

ROOT=Path(__file__).resolve().parent
DATA=json.loads((ROOT/'content.json').read_text())
REPO=ROOT.parent.parent
E=html.escape
ASSETS=ROOT/'assets'
ASSETS.mkdir(exist_ok=True)
for name in ['public-mobile-review.png','public-current.png']:
    shutil.copyfile(REPO/'.artifacts/public-004-acceptance'/name, ASSETS/name)

CSS='''
.claude-viz{--claude-border:rgba(20,20,19,.13)}
.claude-viz{--claude-paper:#FAF9F5;--claude-surface:#FFFDF8;--claude-ink:#141413;--claude-ink-muted:#5F5E58;--claude-mid-gray:#B0AEA5;--claude-mist:#E8E6DC;--claude-orange:#D97757;--claude-orange-soft:#F3DFD5;--claude-blue:#6A9BCC;--claude-blue-soft:#E1EBF4;--claude-green:#788C5D;--claude-green-soft:#E4E9DC;background:var(--claude-paper);color:var(--claude-ink);font:17px/1.75 Lora,Georgia,"Noto Serif SC","Songti SC",serif;margin:0}
.claude-viz *{box-sizing:border-box}.claude-viz h1,.claude-viz h2,.claude-viz h3,.claude-viz nav,.claude-viz table,.claude-viz figcaption,.claude-viz .diagram,.claude-viz .eyebrow{font-family:Poppins,Arial,"PingFang SC","Noto Sans SC",sans-serif}
.claude-viz main{max-width:1200px;margin:auto;padding:56px 48px 90px}.claude-viz h1{font-size:48px;line-height:1.2;letter-spacing:-.025em;max-width:850px;margin:20px 0 24px}.claude-viz h2{font-size:30px;line-height:1.35;margin:0 0 24px}.claude-viz h3{font-size:21px;margin:8px 0 16px}.claude-viz .eyebrow{font-size:13px;letter-spacing:.08em;color:var(--claude-ink-muted)}
.claude-viz .intro{font-size:21px;max-width:850px}.claude-viz nav{display:flex;flex-wrap:wrap;gap:8px;margin:30px 0}.claude-viz a{color:#8E412C;text-underline-offset:4px;overflow-wrap:anywhere}.claude-viz nav a,.claude-viz .links a{display:inline-flex;align-items:center;min-height:44px;padding:8px 14px;border:1px solid var(--claude-mist);border-radius:8px;text-decoration:none;font-size:14px}.claude-viz a:hover{background:var(--claude-orange-soft)}.claude-viz a:focus-visible{outline:2px solid var(--claude-orange);outline-offset:3px}
.claude-viz section{padding:48px 0;border-top:1px solid var(--claude-mist);scroll-margin-top:24px}.claude-viz p{margin:0 0 18px}.claude-viz .lead{font-size:21px;color:var(--claude-ink);max-width:950px}.claude-viz .links{display:flex;flex-wrap:wrap;gap:8px;margin-top:24px}.claude-viz table{width:100%;border-collapse:collapse;margin:24px 0;font-size:14px;line-height:1.6}.claude-viz th,.claude-viz td{padding:13px 15px;border:1px solid var(--claude-mist);text-align:left;vertical-align:top;overflow-wrap:anywhere}.claude-viz th{background:var(--claude-mist);font-weight:600}.claude-viz .sources{font:12px/1.7 Arial,"PingFang SC",sans-serif;color:var(--claude-ink-muted);overflow-wrap:anywhere}.claude-viz .sources a{display:block;min-height:26px}.claude-viz figcaption{color:var(--claude-ink-muted);font-size:13px;margin-top:12px}.claude-viz figure{margin:24px 0}.claude-viz .screen-pair{display:grid;grid-template-columns:280px 1fr;align-items:start;gap:28px}.claude-viz .screen-pair img{max-width:100%;border:1px solid var(--claude-mist);border-radius:8px}.claude-viz .screen-pair .screen-copy{padding:24px 0;font-family:Arial,"PingFang SC",sans-serif}.claude-viz .screen-copy strong{font-size:24px;font-weight:500;display:block;margin-bottom:20px}
.claude-viz .diagram{background:var(--claude-surface);padding:32px;border:1px solid var(--claude-mist);border-radius:14px;line-height:1.45}.claude-viz .diagram-title{font-size:24px;font-weight:600;margin-bottom:8px}.claude-viz .diagram-note{font-size:14px;color:var(--claude-ink-muted);margin-bottom:28px}.claude-viz .flow{display:grid;grid-template-columns:repeat(5,1fr);gap:22px}.claude-viz .node{border:1px solid var(--claude-mist);background:var(--claude-paper);border-radius:10px;padding:20px 14px;position:relative}.claude-viz .node:not(:last-child):after{content:'→';position:absolute;right:-20px;top:36px;color:var(--claude-mid-gray)}.claude-viz .node small{display:block;color:var(--claude-ink-muted);margin-top:12px;font-size:13px}.claude-viz .node b{font-size:17px;display:block}.claude-viz .num{font-size:12px;color:#8E412C;display:block;margin-bottom:14px}.claude-viz .loop-note{border-left:3px solid var(--claude-orange);padding-left:14px;font-size:14px;margin-top:24px}
.claude-viz .lanes{display:grid;grid-template-columns:1.05fr 1fr 1.4fr 1fr;gap:16px}.claude-viz .lane{padding:20px 16px;border-radius:10px;background:var(--claude-paper);border:1px solid var(--claude-mist)}.claude-viz .lane.model{background:var(--claude-orange-soft)}.claude-viz .lane.engine{background:var(--claude-green-soft)}.claude-viz .lane.data{background:var(--claude-blue-soft)}.claude-viz .lane h3{font-size:18px}.claude-viz .lane .step{font-size:14px;padding:14px 0;border-top:1px solid rgba(20,20,19,.13)}.claude-viz .lane .step small{display:block;font-size:12px;margin-top:7px;color:var(--claude-ink-muted)}.claude-viz .boundary{margin-top:20px;padding:16px;background:var(--claude-ink);color:var(--claude-paper);border-radius:8px;font-size:14px}.claude-viz .cache-rows{display:grid;gap:14px}.claude-viz .cache-row{display:grid;grid-template-columns:160px 1fr 220px;align-items:center;gap:20px;border-top:1px solid var(--claude-mist);padding-top:16px;font-size:15px}.claude-viz .cache-row small{font-size:12px;color:var(--claude-ink-muted)}
@media(max-width:720px){.claude-viz main{padding:28px 20px 60px}.claude-viz h1{font-size:34px}.claude-viz h2{font-size:26px}.claude-viz .diagram{padding:22px 18px}.claude-viz .flow,.claude-viz .lanes{grid-template-columns:1fr}.claude-viz .node:not(:last-child):after{content:'↓';right:50%;top:auto;bottom:-23px}.claude-viz .cache-row{grid-template-columns:1fr;gap:8px}.claude-viz .screen-pair{grid-template-columns:1fr}.claude-viz .screen-pair img{max-width:320px}.claude-viz table{font-size:12px}.claude-viz th,.claude-viz td{padding:9px 7px}}
@media(prefers-reduced-motion:reduce){.claude-viz *{scroll-behavior:auto!important;animation:none!important;transition:none!important}}
'''

def diagram(name):
    if name=='journey':
        items=[('表达想法','口语输入／完整示例'),('审阅与编辑','股票、条件、且／或'),('执行与等待','真实状态／有界重试'),('查看证据','收益、基准、成交'),('修改再验证','新版本保留旧报告')]
        nodes=''.join(f'<div class="node"><span class="num">0{i+1}</span><b>{a}</b><small>{b}</small></div>' for i,(a,b) in enumerate(items))
        body=f'<div class="flow">{nodes}</div><div class="loop-note">反馈回路　修改一个参数 → 返回“审阅与编辑” → 生成新运行，不能覆盖旧报告</div>'
        title,note='把一句话，走成一次可核验的实验','用户始终能确认规则、看到进度、追溯结果。'
    elif name=='architecture':
        body='''<div class="lanes"><div class="lane"><h3>01 用户界面</h3><div class="step">React／TypeScript<small>输入 · 联想 · 条件编辑</small></div><div class="step">确认策略版本 →<small>异步受理与状态查询</small></div><div class="step">报告与修改入口<small>数值和模型解释分层</small></div></div><div class="lane model"><h3>02 模型层</h3><div class="step">DeepSeek Flash<small>原话 → 候选JSON</small></div><div class="step">DeepSeek Pro<small>规划 · 缺项补查 · 解读</small></div><div class="step">提案交回校验 →<small>不直接生成收益数</small></div></div><div class="lane engine"><h3>03 工程与 Python</h3><div class="step">能力目录＋原文／参数校验<small>不支持的条件不能绕过</small></div><div class="step">确定性逐日执行<small>信号 → 次日成交尝试 → 净值</small></div><div class="step">事实统计 → 模型解读<small>结果与runId／哈希绑定</small></div></div><div class="lane data"><h3>04 数据层</h3><div class="step">东方财富选股／查数<small>真实序列与来源证据</small></div><div class="step">名称／历史／指标缓存<small>命中复用 · 缺项有限补查</small></div><div class="step">SQLite与运行产物<small>当前为临时实例存储</small></div></div></div><div class="boundary">可信边界　模型候选 → 工程校验 → Python计算 → 已核验事实 → 模型解释</div>'''
        title,note='自然语言与计算之间，保留一道确定性边界','橙色：理解与解释　绿色：校验与计算　蓝色：真实数据与存储；各列为职责，不表示单向串行。'
    else:
        body='''<div class="cache-rows"><div class="cache-row"><b>搜索股票名称</b><span>前端静态名录 → 原位联想</span><small>5567条名称／代码；不含行情</small></div><div class="cache-row"><b>读取日线历史</b><span>内存 → 五股磁盘 → 东方财富</span><small>按股票和区间复用；无统一TTL</small></div><div class="cache-row"><b>读取历史指标</b><span>五股磁盘 → 东方财富字段</span><small>默认24小时；参数与字段匹配</small></div><div class="cache-row"><b>要求重新读取</b><span>跳过普通缓存 → 独立实时取数</span><small>避免旧在途结果覆盖新请求</small></div></div><div class="loop-note">不变原则　输入缓存命中后仍由 Python 按当前策略计算；来源、口径与日期仍须校验。</div>'''
        title,note='缓存复用输入，不替代策略计算','名称、行情、指标和报告是四种不同对象，不能混称“全A股离线缓存”。'
    return f'<figure class="diagram" id="fig-{name}"><div class="diagram-title">{title}</div><div class="diagram-note">{note}</div>{body}</figure>'

def table(rows):
    if not rows:return ''
    return '<table><thead><tr>'+''.join(f'<th>{E(x)}</th>' for x in rows[0])+'</tr></thead><tbody>'+''.join('<tr>'+''.join(f'<td>{E(x)}</td>' for x in row)+'</tr>' for row in rows[1:])+'</tbody></table>'

def source_url(path):
    if path.startswith('docs/release-public-005') or path.startswith('docs/releases/public-005'):return '../'+path.removeprefix('docs/')
    if path=='docs/bugfix-concurrency-20260906.md':return '../bugfix-concurrency-20260906.md'
    if path=='docs/bug-backlog-and-acceptance.md':return '../bug-backlog-and-acceptance.md'
    return 'https://github.com/kasaviawuh111-cpu/ai-strategy-backtest/blob/19220c1/'+path

def section(s,details=False):
    out=f'<section id="{s["id"]}"><h2>{E(s["title"])}</h2>'
    if s.get('lead'):out+=f'<p class="lead">{E(s["lead"])}</p>'
    if s.get('figure'):out+=diagram(s['figure'])
    out+=''.join('<p>'+E(p)+'</p>' for p in s.get('paragraphs',[]))
    out+=table(s.get('rows'))
    out+=''.join('<p>'+E(p)+'</p>' for p in s.get('paragraphsAfter',[]))
    if s.get('screenshot'):out+=f'<figure class="screen-pair"><img src="assets/{s["screenshot"]}" alt="004公网实际条件审阅页面"><figcaption class="screen-copy"><strong>从整句话，变成可编辑的条件。</strong><p>条件逐项展示；全部满足与任意触发明确可见；股票与区间可直接修改。</p>{E(s["caption"])}</figcaption></figure>'
    if s.get('links'):out+='<div class="links">'+''.join(f'<a href="#{a}">{b} ↗</a>' for a,b in s['links'])+'</div>'
    if s.get('sources'):out+='<div class="sources">核对依据（源码链接需要私有仓库访问权限）'+''.join(f'<a href="{source_url(p)}">{E(p)}</a>' for p in s['sources'])+'</div>'
    return out+'</section>'

def build_html():
    title=DATA['title']
    nav=''.join(f'<a href="#{s["id"]}">{s["title"]}</a>' for s in DATA['appendix'])
    page=f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>{title}</title><style>{CSS}</style><body class="claude-viz"><main><div class="eyebrow">PRODUCT & ENGINEERING REVIEW · 2026.09.06 · PUBLIC 005</div><h1>让投资想法<br>变成可核对的历史实验</h1><p class="intro">A股自然语言策略回测阶段汇报<br>领导速读正文＋可点击技术与验收附录</p><nav><a href="#summary">阅读结论</a><a href="#architecture">系统架构</a><a href="#evidence">验收证据</a><a href="{DATA["publicUrl"]}">打开公网产品 ↗</a></nav>'+''.join(section(s) for s in DATA['brief'])+f'<section id="details"><div class="eyebrow">TECHNICAL APPENDIX</div><h2>按需展开的工程事实</h2><p>截至2026年9月6日。已实现、配置核验、真实公网验收与未完成项分别表述。完整源码仍受私有仓库权限控制。</p><nav>{nav}</nav></section>'+''.join(section(s,True) for s in DATA['appendix'])+'</main></body></html>'
    (ROOT/'index.html').write_text(page)
    (ROOT/'figures.html').write_text(f'<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>架构与流程示意</title><style>{CSS}</style><body class="claude-viz"><main>'+''.join(diagram(k) for k in ['journey','architecture','cache'])+'</main></body></html>')
    for name,seq in [('report',DATA['brief']),('technical-appendix',DATA['appendix'])]:
        md=['# '+(DATA['title'] if name=='report' else 'A股策略回测技术与验收附录'),'截至2026年9月6日 · 公网005']
        for s in seq:
            md+=['\n## '+s['title'],s.get('lead',''),*s.get('paragraphs',[])]
            if s.get('rows'):
                rows=s['rows'];md+=['\n'.join(['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']+['| '+' | '.join(r)+' |' for r in rows[1:]])]
            md+=s.get('paragraphsAfter',[])
            md+=['['+b+'](index.html#'+a+')' for a,b in s.get('links',[])]
            md+=['依据 '+p for p in s.get('sources',[])]
        (ROOT/(name+'.md')).write_text('\n\n'.join(md))

def build_docx():
    from docx import Document
    from docx.shared import Inches,Pt,RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    urls=json.loads((ROOT/'feishu-links.json').read_text()) if (ROOT/'feishu-links.json').exists() else {}
    def link(p,text,url,anchor=None):
        h=OxmlElement('w:hyperlink')
        if anchor:h.set(qn('w:anchor'),anchor)
        else:h.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'8E412C');pr.append(col);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)
    def bookmark(p,name,i):
        a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(i));a.set(qn('w:name'),name);p._p.insert(0,a)
        b=OxmlElement('w:bookmarkEnd');b.set(qn('w:id'),str(i));p._p.append(b)
    def setup(doc,label):
        sec=doc.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
        sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
        for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
            st=doc.styles[name];st.font.name='Arial';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Heiti SC');st.font.color.rgb=RGBColor.from_string('141413')
        for border in doc.styles.element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
        for fonts in doc.styles.element.xpath('.//w:rFonts'):
            fonts.set(qn('w:eastAsia'),'Heiti SC')
            for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme']:
                if qn('w:'+key) in fonts.attrib:del fonts.attrib[qn('w:'+key)]
        st=doc.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.line_spacing=1.25;st.paragraph_format.space_after=Pt(8)
        for name,size in [('Title',25),('Heading 1',19),('Heading 2',14)]:
            st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_after=Pt(12)
        doc.styles['Caption'].font.size=Pt(9)
        h=sec.header.paragraphs[0];h.text=label+'  ·  公网005  ·  2026.09.06';h.style='Caption'
        foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        foot.add_run('阶段汇报  |  ')
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
        doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=label
    def addtable(doc,rows):
        tb=doc.add_table(rows=0,cols=len(rows[0]));tb.autofit=False
        for ri,row in enumerate(rows):
            cells=tb.add_row().cells
            for c,txt in zip(cells,row):
                c.text=txt
                for p in c.paragraphs:
                    p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.12
                    for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
                pr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    item=OxmlElement('w:'+edge);item.set(qn('w:w'),'90');item.set(qn('w:type'),'dxa');m.append(item)
                pr.append(m)
            trpr=tb.rows[-1]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
            if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        b=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
        tb._tbl.tblPr.append(b)
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
    def image(doc,path,width,alt):
        pic=doc.add_picture(str(path),width=Inches(width));pic._inline.docPr.set('descr',alt)
    for kind,seq in [('brief',DATA['brief']),('technical-appendix',DATA['appendix'])]:
        doc=Document();setup(doc,DATA['title'] if kind=='brief' else 'A股策略回测技术与验收附录')
        for i,s in enumerate(seq):
            if i:doc.add_page_break()
            if i==0:
                doc.add_heading(DATA['title'] if kind=='brief' else 'A股策略回测技术与验收附录',0)
                p=doc.add_paragraph();link(p,'公网产品',DATA['publicUrl']);p.add_run('  ·  核验基准 2026年9月6日')
            h=doc.add_heading(s['title'].replace('，',' ').replace('／',' '),1);bookmark(h,s['id'],i+1)
            if s.get('lead'):
                p=doc.add_paragraph(s['lead']);p.runs[0].bold=True
            if s.get('figure') and not (kind=='technical-appendix' and s['figure']=='cache'):
                image(doc,ASSETS/(s['figure']+'.png'),6.75,s['title']+'示意图')
            for p in s.get('paragraphs',[]):doc.add_paragraph(p)
            if s.get('rows'):addtable(doc,s['rows'])
            for p in s.get('paragraphsAfter',[]):doc.add_paragraph(p)
            if s.get('screenshot'):
                # A compact real UI image rather than a full-width phone screenshot.
                image(doc,ASSETS/s['screenshot'],1.75,'004公网审阅页，参数修改前RSI30')
                doc.add_paragraph(s['caption'],'Caption')
            if s.get('links'):
                p=doc.add_paragraph()
                for n,(a,b) in enumerate(s['links']):
                    if n:p.add_run('  ·  ')
                    if a in [x['id'] for x in DATA['brief']]:link(p,b,'',a)
                    elif urls.get('appendix'):link(p,b+'（附录'+str(next(j+1 for j,x in enumerate(DATA['appendix']) if x['id']==a))+'）',urls['appendix'])
                    else:link(p,b,'technical-appendix.docx',None)
            if s.get('sources'):
                p=doc.add_paragraph('核对依据（源码链接需私有仓库权限）','Caption')
                for src in s['sources']:
                    p=doc.add_paragraph(style='Caption')
                    if src.endswith('bug-backlog-and-acceptance.md'):p.add_run(src+'，当前剩余待办与归档依据')
                    else:link(p,src,source_url(src))
        doc.save(ROOT/(kind+'.docx'))

if __name__=='__main__':
    if '--docx' in sys.argv:build_docx()
    else:build_html()
